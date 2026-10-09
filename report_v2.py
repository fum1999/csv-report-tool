import pandas as pd
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


# ============================================================
# 設定
# ============================================================

INPUT_FILE = "sample_sales.csv"
OUTPUT_FILE = "sales_report.xlsx"

# Excel用カラー
NAVY = "1F4E78"
BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F3F6F8"
WHITE = "FFFFFF"
DARK = "1F2937"
GREEN = "70AD47"
ORANGE = "ED7D31"

THIN_BORDER = Side(
    style="thin",
    color="D9E1F2"
)


# ============================================================
# 1. CSV読み込み
# ============================================================

print("CSVを読み込んでいます...")

df = pd.read_csv(INPUT_FILE)

# 必須列チェック
required_columns = ["日付", "商品", "売上"]

missing_columns = [
    col for col in required_columns
    if col not in df.columns
]

if missing_columns:
    raise ValueError(
        f"CSVに必要な列がありません: {missing_columns}"
    )

# データ型を整理
df["日付"] = pd.to_datetime(
    df["日付"],
    errors="coerce"
)

df["売上"] = pd.to_numeric(
    df["売上"],
    errors="coerce"
)

# 不正データチェック
if df["日付"].isna().any():
    raise ValueError("日付に読み込めないデータがあります。")

if df["売上"].isna().any():
    raise ValueError("売上に数値として読み込めないデータがあります。")

# 日付順に並べる
df = df.sort_values("日付")


# ============================================================
# 2. 集計
# ============================================================

print("売上を集計しています...")

# 商品別
product_sales = (
    df.groupby("商品", as_index=False)["売上"]
    .sum()
    .sort_values(
        "売上",
        ascending=False
    )
)

# 日別
daily_sales = (
    df.groupby("日付", as_index=False)["売上"]
    .sum()
    .sort_values("日付")
)

# 総売上
total_sales = df["売上"].sum()

# 平均日次売上
average_daily_sales = daily_sales["売上"].mean()

# 売上トップ商品
best_product_row = product_sales.iloc[0]

best_product = best_product_row["商品"]
best_product_sales = best_product_row["売上"]

# 売上最高日
best_day_row = daily_sales.loc[
    daily_sales["売上"].idxmax()
]

best_day = best_day_row["日付"]
best_day_sales = best_day_row["売上"]

# 集計期間
start_date = df["日付"].min()
end_date = df["日付"].max()

# 商品数
product_count = df["商品"].nunique()

# データ件数
data_count = len(df)


# ============================================================
# 3. Excelブック作成
# ============================================================

print("Excelレポートを作成しています...")

wb = Workbook()

# デフォルトシート
ws_summary = wb.active
ws_summary.title = "概要"

# その他シート
ws_product = wb.create_sheet("商品別売上")
ws_daily = wb.create_sheet("日別売上")
ws_raw = wb.create_sheet("元データ")


# ============================================================
# 4. 共通関数
# ============================================================

def set_title(ws, title, subtitle=None):

    ws.merge_cells("A1:H1")

    cell = ws["A1"]

    cell.value = title
    cell.font = Font(
        size=22,
        bold=True,
        color=WHITE
    )

    cell.fill = PatternFill(
        "solid",
        fgColor=NAVY
    )

    cell.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    ws.row_dimensions[1].height = 36

    if subtitle:

        ws.merge_cells("A2:H2")

        cell = ws["A2"]

        cell.value = subtitle

        cell.font = Font(
            size=10,
            color="666666"
        )

        cell.alignment = Alignment(
            vertical="center"
        )

        ws.row_dimensions[2].height = 22


def style_header(ws, row, start_col, end_col):

    for col in range(start_col, end_col + 1):

        cell = ws.cell(
            row=row,
            column=col
        )

        cell.font = Font(
            bold=True,
            color=WHITE
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=NAVY
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.border = Border(
            bottom=THIN_BORDER
        )


def auto_width(ws):

    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is None:
                continue

            length = len(str(cell.value))

            if length > max_length:
                max_length = length

        ws.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 3, 12),
            30
        )


def setup_sheet(ws):

    ws.sheet_view.showGridLines = False

    ws.freeze_panes = "A4"


# ============================================================
# 5. 概要シート
# ============================================================

set_title(
    ws_summary,
    "売上レポート",
    f"集計期間：{start_date:%Y/%m/%d} ～ {end_date:%Y/%m/%d}"
)

ws_summary["A4"] = "主要KPI"

ws_summary["A4"].font = Font(
    size=14,
    bold=True,
    color=DARK
)


# KPI
kpis = [

    (
        "総売上",
        total_sales,
        '¥#,##0'
    ),

    (
        "1日平均売上",
        average_daily_sales,
        '¥#,##0'
    ),

    (
        "売上トップ商品",
        best_product,
        '@'
    ),

    (
        "トップ商品売上",
        best_product_sales,
        '¥#,##0'
    )

]


for i, (label, value, number_format) in enumerate(kpis):

    start_col = 1 + i * 2

    # ラベル
    ws_summary.merge_cells(
        start_row=5,
        start_column=start_col,
        end_row=5,
        end_column=start_col + 1
    )

    label_cell = ws_summary.cell(
        5,
        start_col
    )

    label_cell.value = label

    label_cell.font = Font(
        bold=True,
        color=NAVY
    )

    label_cell.fill = PatternFill(
        "solid",
        fgColor=LIGHT_BLUE
    )

    label_cell.alignment = Alignment(
        horizontal="center"
    )

    # 値
    ws_summary.merge_cells(
        start_row=6,
        start_column=start_col,
        end_row=6,
        end_column=start_col + 1
    )

    value_cell = ws_summary.cell(
        6,
        start_col
    )

    value_cell.value = value

    value_cell.font = Font(
        size=16,
        bold=True,
        color=DARK
    )

    value_cell.alignment = Alignment(
        horizontal="center"
    )

    value_cell.number_format = number_format

    # 枠線
    for row in [5, 6]:

        for col in range(
            start_col,
            start_col + 2
        ):

            ws_summary.cell(
                row,
                col
            ).border = Border(
                left=THIN_BORDER,
                right=THIN_BORDER,
                top=THIN_BORDER,
                bottom=THIN_BORDER
            )


# ============================================================
# 6. 概要シートの商品ランキング
# ============================================================

ws_summary["A9"] = "商品別売上ランキング"

ws_summary["A9"].font = Font(
    size=14,
    bold=True,
    color=DARK
)

ws_summary["A10"] = "商品"
ws_summary["B10"] = "売上"

style_header(
    ws_summary,
    10,
    1,
    2
)


for row_num, row in enumerate(
    product_sales.itertuples(index=False),
    start=11
):

    ws_summary.cell(
        row_num,
        1,
        row.商品
    )

    ws_summary.cell(
        row_num,
        2,
        row.売上
    )

    ws_summary.cell(
        row_num,
        2
    ).number_format = '¥#,##0'


# ============================================================
# 7. 商品別売上シート
# ============================================================

set_title(
    ws_product,
    "商品別売上",
    "商品ごとの売上集計"
)

ws_product["A4"] = "商品"
ws_product["B4"] = "売上"

style_header(
    ws_product,
    4,
    1,
    2
)


for row_num, row in enumerate(
    product_sales.itertuples(index=False),
    start=5
):

    ws_product.cell(
        row_num,
        1,
        row.商品
    )

    ws_product.cell(
        row_num,
        2,
        row.売上
    )

    ws_product.cell(
        row_num,
        2
    ).number_format = '¥#,##0'


# 条件付き書式
last_product_row = 4 + len(product_sales)

ws_product.conditional_formatting.add(

    f"B5:B{last_product_row}",

    ColorScaleRule(
        start_type="min",
        start_color="EAF2F8",

        mid_type="percentile",
        mid_value=50,
        mid_color="AED6F1",

        end_type="max",
        end_color="2E86C1"
    )

)


# ============================================================
# 8. 商品別棒グラフ
# ============================================================

chart_product = BarChart()

chart_product.type = "col"

chart_product.style = 10

chart_product.title = "商品別売上"

chart_product.y_axis.title = "売上（円）"

chart_product.x_axis.title = "商品"

data_ref = Reference(
    ws_product,
    min_col=2,
    min_row=4,
    max_row=last_product_row
)

category_ref = Reference(
    ws_product,
    min_col=1,
    min_row=5,
    max_row=last_product_row
)

chart_product.add_data(
    data_ref,
    titles_from_data=True
)

chart_product.set_categories(
    category_ref
)

chart_product.height = 8
chart_product.width = 14

chart_product.legend = None

chart_product.dataLabels = DataLabelList()
chart_product.dataLabels.showVal = True

ws_product.add_chart(
    chart_product,
    "D4"
)


# ============================================================
# 9. 日別売上シート
# ============================================================

set_title(
    ws_daily,
    "日別売上",
    "日ごとの売上推移"
)

ws_daily["A4"] = "日付"
ws_daily["B4"] = "売上"

style_header(
    ws_daily,
    4,
    1,
    2
)


for row_num, row in enumerate(
    daily_sales.itertuples(index=False),
    start=5
):

    ws_daily.cell(
        row_num,
        1,
        row.日付.to_pydatetime()
    )

    ws_daily.cell(
        row_num,
        1
    ).number_format = "yyyy/mm/dd"

    ws_daily.cell(
        row_num,
        2,
        row.売上
    )

    ws_daily.cell(
        row_num,
        2
    ).number_format = '¥#,##0'


last_daily_row = 4 + len(daily_sales)


# ============================================================
# 10. 日別折れ線グラフ
# ============================================================

chart_daily = LineChart()

chart_daily.style = 13

chart_daily.title = "日別売上推移"

chart_daily.y_axis.title = "売上（円）"

chart_daily.x_axis.title = "日付"

data_ref = Reference(
    ws_daily,
    min_col=2,
    min_row=4,
    max_row=last_daily_row
)

category_ref = Reference(
    ws_daily,
    min_col=1,
    min_row=5,
    max_row=last_daily_row
)

chart_daily.add_data(
    data_ref,
    titles_from_data=True
)

chart_daily.set_categories(
    category_ref
)

chart_daily.height = 8
chart_daily.width = 15

chart_daily.legend = None

ws_daily.add_chart(
    chart_daily,
    "D4"
)


# ============================================================
# 11. 元データシート
# ============================================================

set_title(
    ws_raw,
    "元データ",
    "レポート作成に使用した明細データ"
)

headers = [
    "日付",
    "商品",
    "売上"
]

for col_num, header in enumerate(
    headers,
    start=1
):

    ws_raw.cell(
        4,
        col_num,
        header
    )

style_header(
    ws_raw,
    4,
    1,
    3
)


for row_num, row in enumerate(
    df.itertuples(index=False),
    start=5
):

    ws_raw.cell(
        row_num,
        1,
        row.日付.to_pydatetime()
    )

    ws_raw.cell(
        row_num,
        1
    ).number_format = "yyyy/mm/dd"

    ws_raw.cell(
        row_num,
        2,
        row.商品
    )

    ws_raw.cell(
        row_num,
        3,
        row.売上
    )

    ws_raw.cell(
        row_num,
        3
    ).number_format = '¥#,##0'


# ============================================================
# 12. 元データをExcelテーブル化
# ============================================================

last_raw_row = 4 + len(df)

table = Table(
    displayName="SalesData",
    ref=f"A4:C{last_raw_row}"
)

table_style = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)

table.tableStyleInfo = table_style

ws_raw.add_table(table)


# ============================================================
# 13. シート共通設定
# ============================================================

for ws in wb.worksheets:

    setup_sheet(ws)

    auto_width(ws)

    # 行の高さ
    ws.row_dimensions[4].height = 25

    # 印刷設定
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.sheet_properties.pageSetUpPr.fitToPage = True


# 概要シートだけ固定幅を調整
ws_summary.column_dimensions["A"].width = 22
ws_summary.column_dimensions["B"].width = 18
ws_summary.column_dimensions["C"].width = 22
ws_summary.column_dimensions["D"].width = 18
ws_summary.column_dimensions["E"].width = 22
ws_summary.column_dimensions["F"].width = 18
ws_summary.column_dimensions["G"].width = 22
ws_summary.column_dimensions["H"].width = 18


# ============================================================
# 14. 概要に追加情報
# ============================================================

ws_summary["A16"] = "売上ピーク日"

ws_summary["B16"] = best_day.to_pydatetime()

ws_summary["B16"].number_format = "yyyy/mm/dd"

ws_summary["C16"] = best_day_sales

ws_summary["C16"].number_format = '¥#,##0'


ws_summary["E16"] = "商品数"

ws_summary["F16"] = product_count

ws_summary["G16"] = "明細件数"

ws_summary["H16"] = data_count


for cell in [
    "A16",
    "B16",
    "C16",
    "E16",
    "F16",
    "G16",
    "H16"
]:

    ws_summary[cell].border = Border(
        left=THIN_BORDER,
        right=THIN_BORDER,
        top=THIN_BORDER,
        bottom=THIN_BORDER
    )


ws_summary["A16"].font = Font(
    bold=True
)

ws_summary["E16"].font = Font(
    bold=True
)

ws_summary["G16"].font = Font(
    bold=True
)


# ============================================================
# 15. アクティブシート
# ============================================================

wb.active = 0


# ============================================================
# 16. 保存
# ============================================================

wb.save(OUTPUT_FILE)


# ============================================================
# 17. 完了メッセージ
# ============================================================

print()
print("=" * 50)
print("売上レポート作成完了")
print("=" * 50)

print(f"出力ファイル : {OUTPUT_FILE}")
print(f"総売上       : ¥{total_sales:,.0f}")
print(f"1日平均売上  : ¥{average_daily_sales:,.0f}")
print(f"トップ商品   : {best_product}")
print(f"トップ商品売上: ¥{best_product_sales:,.0f}")
print(
    f"売上ピーク日 : {best_day:%Y/%m/%d}"
    f"（¥{best_day_sales:,.0f}）"
)
print(f"商品数       : {product_count}")
print(f"明細件数     : {data_count}")
print("=" * 50)

