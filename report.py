import pandas as pd
import matplotlib.pyplot as plt
plt.rcParams["font.family"] = "Meiryo"

# CSV読込
df = pd.read_csv("sample_sales.csv")

# 商品別売上集計
product_sales = (
    df.groupby("商品")["売上"]
    .sum()
    .reset_index()
)

# 日別売上集計
daily_sales = (
    df.groupby("日付")["売上"]
    .sum()
    .reset_index()
)

# 総売上
total_sales = df["売上"].sum()

print("総売上:", total_sales)

# グラフ作成
plt.figure(figsize=(8, 5))

plt.bar(
    product_sales["商品"],
    product_sales["売上"]
)

plt.title("商品別売上")
plt.xlabel("商品")
plt.ylabel("売上")

plt.tight_layout()

plt.savefig("sales_chart.png")

# Excel出力
with pd.ExcelWriter(
    "report.xlsx",
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="元データ",
        index=False
    )

    daily_sales.to_excel(
        writer,
        sheet_name="日別売上",
        index=False
    )

    product_sales.to_excel(
        writer,
        sheet_name="商品別売上",
        index=False
    )

print("レポート作成完了")

from openpyxl import load_workbook
from openpyxl.drawing.image import Image

wb = load_workbook("report.xlsx")

ws = wb.create_sheet("グラフ")

img = Image("sales_chart.png")

ws.add_image(img, "A1")

wb.save("report.xlsx")